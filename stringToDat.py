import os
import sys
from collections import defaultdict
from openpyxl import load_workbook

WORKING_DIR = os.path.dirname(os.path.realpath(__file__))
IN_FILE = os.path.join(WORKING_DIR, "ModLanguages.xlsx")
OUT_DIR = os.path.join(WORKING_DIR, "TownOfWeekend", "Resources", "Languages")

def excel_to_dat(in_file):
    wb = load_workbook(in_file, read_only=True)
    language_data = defaultdict(dict)

    for sheet in wb:
        sheet_name = sheet.title
        headers = [cell.value for cell in sheet[1][1:]]  # 第一行从第二列开始为文件名

        for row in sheet.iter_rows(min_row=2):
            key_base = row[0].value
            if not key_base:
                continue

            # 获取第二列的默认值（如果后续列内容为空则用此值）
            default_value = row[1].value if row[1].value is not None else ""

            for idx, cell in enumerate(row[1:]):  # 遍历语言列
                if idx >= len(headers) or not headers[idx]:
                    continue

                filename = f"{headers[idx]}.dat"
                full_key = f"{sheet_name}.{key_base}"

                # 若当前单元格为空，则使用第二列的值
                current_value = cell.value if cell.value is not None else ""
                if not str(current_value).strip():
                    current_value = default_value

                # 处理换行（保留Excel中的换行）
                value = str(current_value).replace('\r\n', '\n').replace('\r', '\n')
                language_data[filename][full_key] = value

    os.makedirs(OUT_DIR, exist_ok=True)
    for filename, data in language_data.items():
        output_path = os.path.join(OUT_DIR, filename)
        with open(output_path, 'w', encoding='utf-8') as f:
            for key, value in sorted(data.items()):
                f.write(f'"{key}": "{value}"\n')

if __name__ == "__main__":
    if not os.path.exists(IN_FILE):
        print(f"错误：找不到输入文件 {IN_FILE}")
        sys.exit(1)
    
    excel_to_dat(IN_FILE)
    print(f"转换完成！文件已输出至 {OUT_DIR}")